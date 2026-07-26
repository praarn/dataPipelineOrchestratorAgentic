"""
Renders the Report Agent's markdown into downloadable PDF / DOCX files.

This is intentionally a small, purpose-built parser rather than a generic
markdown engine: `report_agent.build_report` only ever emits a known set of
constructs (headings, bullets, blockquotes, embedded base64 images, simple
pipe tables, italic captions, horizontal rules), so a full CommonMark parser
would be overkill. If the report template grows more elaborate constructs
later, extend `_parse_blocks` rather than reaching for a markdown library.
"""
from __future__ import annotations

import base64
import io
import re


def _parse_blocks(markdown: str) -> list[dict]:
    blocks: list[dict] = []
    lines = markdown.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        if line.strip() == "---":
            blocks.append({"type": "hr"})
            i += 1
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading_match:
            blocks.append({"type": "heading", "level": len(heading_match.group(1)),
                            "text": heading_match.group(2).strip()})
            i += 1
            continue

        img_match = re.match(r"^!\[[^\]]*\]\(data:image/png;base64,([A-Za-z0-9+/=]+)\)", line)
        if img_match:
            try:
                blocks.append({"type": "image", "data": base64.b64decode(img_match.group(1))})
            except Exception:
                pass
            i += 1
            continue

        if line.startswith("> "):
            blocks.append({"type": "quote", "text": line[2:].strip()})
            i += 1
            continue

        if line.startswith("- "):
            blocks.append({"type": "bullet", "text": line[2:].strip()})
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if all(re.match(r"^:?-+:?$", c) for c in cells):
                    continue  # markdown separator row
                rows.append(cells)
            if rows:
                blocks.append({"type": "table", "rows": rows})
            continue

        if line.strip().startswith("*") and line.strip().endswith("*") and len(line.strip()) > 1:
            blocks.append({"type": "italic", "text": line.strip().strip("*")})
            i += 1
            continue

        blocks.append({"type": "para", "text": line.strip()})
        i += 1

    return blocks


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_CODE_RE = re.compile(r"`([^`]+)`")


def _strip_inline_markup(text: str) -> str:
    text = _BOLD_RE.sub(r"\1", text)
    text = _CODE_RE.sub(r"\1", text)
    return text


def markdown_to_pdf_bytes(markdown: str) -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table,
        TableStyle, HRFlowable,
    )

    blocks = _parse_blocks(markdown)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, topMargin=0.75 * inch,
                             bottomMargin=0.75 * inch, leftMargin=0.75 * inch,
                             rightMargin=0.75 * inch)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="H1c", parent=styles["Heading1"], spaceAfter=10))
    styles.add(ParagraphStyle(name="H2c", parent=styles["Heading2"], spaceBefore=14, spaceAfter=8))
    styles.add(ParagraphStyle(name="H3c", parent=styles["Heading3"], spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle(name="Bodyc", parent=styles["BodyText"], spaceAfter=6, leading=15))
    styles.add(ParagraphStyle(name="Bulletc", parent=styles["BodyText"], leftIndent=16,
                               spaceAfter=4, bulletIndent=4, leading=14))
    styles.add(ParagraphStyle(name="Quotec", parent=styles["BodyText"], leftIndent=16,
                               textColor=colors.HexColor("#555555"), spaceAfter=6,
                               borderColor=colors.HexColor("#cccccc"), borderWidth=0,
                               leading=14))
    styles.add(ParagraphStyle(name="Italicc", parent=styles["BodyText"], fontName="Helvetica-Oblique",
                               textColor=colors.HexColor("#666666"), spaceAfter=10, fontSize=9))

    story = []
    heading_style = {1: "H1c", 2: "H2c", 3: "H3c", 4: "H3c"}
    for b in blocks:
        t = b["type"]
        if t == "heading":
            story.append(Paragraph(_strip_inline_markup(b["text"]), styles[heading_style.get(b["level"], "H3c")]))
        elif t == "para":
            story.append(Paragraph(_strip_inline_markup(b["text"]), styles["Bodyc"]))
        elif t == "bullet":
            story.append(Paragraph("• " + _strip_inline_markup(b["text"]), styles["Bulletc"]))
        elif t == "quote":
            story.append(Paragraph(_strip_inline_markup(b["text"]), styles["Quotec"]))
        elif t == "italic":
            story.append(Paragraph(_strip_inline_markup(b["text"]), styles["Italicc"]))
        elif t == "hr":
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", color=colors.HexColor("#dddddd")))
            story.append(Spacer(1, 8))
        elif t == "table":
            rows = b["rows"]
            wrapped = [[Paragraph(_strip_inline_markup(c), styles["Bodyc"]) for c in row] for row in rows]
            tbl = Table(wrapped, hAlign="LEFT")
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f6")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 8))
        elif t == "image":
            try:
                img_buf = io.BytesIO(b["data"])
                img = RLImage(img_buf, width=5.5 * inch, height=3.2 * inch, kind="proportional")
                story.append(img)
                story.append(Spacer(1, 6))
            except Exception:
                continue

    doc.build(story)
    return buf.getvalue()


def markdown_to_docx_bytes(markdown: str) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    blocks = _parse_blocks(markdown)
    doc = Document()

    heading_level = {1: 0, 2: 1, 3: 2, 4: 3}
    for b in blocks:
        t = b["type"]
        if t == "heading":
            level = heading_level.get(b["level"], 2)
            doc.add_heading(_strip_inline_markup(b["text"]), level=level)
        elif t == "para":
            doc.add_paragraph(_strip_inline_markup(b["text"]))
        elif t == "bullet":
            doc.add_paragraph(_strip_inline_markup(b["text"]), style="List Bullet")
        elif t == "quote":
            p = doc.add_paragraph()
            run = p.add_run(_strip_inline_markup(b["text"]))
            run.italic = True
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            p.paragraph_format.left_indent = Inches(0.25)
        elif t == "italic":
            p = doc.add_paragraph()
            run = p.add_run(_strip_inline_markup(b["text"]))
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        elif t == "hr":
            p = doc.add_paragraph("_" * 40)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == "table":
            rows = b["rows"]
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = _strip_inline_markup(cell)
            doc.add_paragraph("")
        elif t == "image":
            try:
                img_buf = io.BytesIO(b["data"])
                doc.add_picture(img_buf, width=Inches(5.8))
            except Exception:
                continue

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
